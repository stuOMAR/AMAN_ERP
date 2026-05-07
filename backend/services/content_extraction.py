"""
T18 P1 #96 — Best-effort full-text extraction from uploaded attachments.

Used by the scheduler job ``extract_attachment_content`` and exposed
through ``GET /attachments/search?q=...`` so users can find documents by
their textual content (contracts, memos, scanned-and-OCRd PDFs, Excel
sheets, etc.).

Design constraints:
  * Library imports are LAZY — we don't want optional deps to crash app
    startup if the operator chose not to install ``pypdf`` / ``openpyxl``
    / ``docx2txt`` on their image.
  * Every extractor swallows its own exceptions and returns ``None`` —
    the scheduler then records ``content_extraction_error`` and moves on
    rather than retrying forever.
  * Output is capped at 200 KB per attachment to keep the GIN index
    size sane on tenants with very large PDFs.
"""
from __future__ import annotations

import logging
import os
from typing import Optional

logger = logging.getLogger(__name__)

# Cap per-attachment extracted text to keep the GIN index manageable.
_MAX_TEXT_BYTES = 200 * 1024  # 200 KB


def _truncate(text: str) -> str:
    if not text:
        return ""
    encoded = text.encode("utf-8", errors="ignore")
    if len(encoded) <= _MAX_TEXT_BYTES:
        return text
    return encoded[:_MAX_TEXT_BYTES].decode("utf-8", errors="ignore")


def _extract_pdf(path: str) -> Optional[str]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return None
    try:
        reader = PdfReader(path)
        chunks = []
        for page in reader.pages:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(chunks).strip() or None
    except Exception:
        return None


def _extract_xlsx(path: str) -> Optional[str]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return None
    try:
        wb = load_workbook(filename=path, read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    chunks.append(" ".join(cells))
        return "\n".join(chunks).strip() or None
    except Exception:
        return None


def _extract_docx(path: str) -> Optional[str]:
    try:
        import docx2txt  # type: ignore
    except Exception:
        return None
    try:
        return (docx2txt.process(path) or "").strip() or None
    except Exception:
        return None


def _extract_text(path: str) -> Optional[str]:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            return fh.read().strip() or None
    except Exception:
        return None


def extract_text(path: str, mime: Optional[str] = None,
                 file_name: Optional[str] = None) -> Optional[str]:
    """Best-effort text extraction. Returns ``None`` on failure (caller
    should record this in ``attachments.content_extraction_error``)."""
    if not path or not os.path.isfile(path):
        return None

    mime = (mime or "").lower()
    name = (file_name or path).lower()

    if mime == "application/pdf" or name.endswith(".pdf"):
        result = _extract_pdf(path)
    elif mime in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",) \
            or name.endswith(".xlsx") or name.endswith(".xlsm"):
        result = _extract_xlsx(path)
    elif mime in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) \
            or name.endswith(".docx"):
        result = _extract_docx(path)
    elif mime.startswith("text/") or name.endswith((".txt", ".csv", ".md", ".log")):
        result = _extract_text(path)
    else:
        return None

    if not result:
        return None
    return _truncate(result)
